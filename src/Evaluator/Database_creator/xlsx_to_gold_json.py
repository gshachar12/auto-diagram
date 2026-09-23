"""
Converts a filled-in gold_sample_template.xlsx into the gold_sample.json
shape run_evaluator.py expects. The example row (row 2 on each sheet) is
always skipped by FIXED POSITION, not by matching its content -- an
earlier version matched on the example's id values ("kali_entity", "c1"),
which would silently drop a real user's own data if they happened to
reuse a similarly plausible id (confirmed directly: this happened
immediately when testing with realistic values).

Usage:
    python3 xlsx_to_gold_json.py filled_template.xlsx -o gold_sample.json
"""
from __future__ import annotations

import argparse
import json

import openpyxl


def convert(xlsx_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    ws_info = wb["Info"]
    sample_id = ws_info["B1"].value
    if not sample_id:
        raise ValueError("Info sheet: sample_id (cell B1) is empty -- fill it in first.")

    ws_a = wb["Actors"]
    actors = []
    for row in ws_a.iter_rows(min_row=3, values_only=True):  # row 2 is the fixed example row, always skipped
        entity_id, description_ref_gold, ground_truth_ip = (row + (None, None, None))[:3]
        if not entity_id:
            continue
        actors.append({
            "entity_id": str(entity_id).strip(),
            "description_ref_gold": str(description_ref_gold or "").strip(),
            "ground_truth_ip": str(ground_truth_ip).strip() if ground_truth_ip else None,
        })

    ws_c = wb["Claims"]
    claims = []
    for row in ws_c.iter_rows(min_row=3, values_only=True):  # row 2 is the fixed example row, always skipped
        (claim_id, text, severity, references_entities,
         gold_grounding_status, gold_linked_fingerprint_keys) = (row + (None,) * 6)[:6]
        if not claim_id:
            continue
        refs = [e.strip() for e in str(references_entities or "").split(",") if e.strip()]
        # Fingerprint keys are themselves JSON-array-shaped strings
        # containing commas (e.g. ["ARP",["2","aa:bb"]]) -- a naive split
        # on "," would break them apart. Split only on a comma that is
        # NOT inside brackets.
        raw_keys = str(gold_linked_fingerprint_keys or "").strip()
        keys = _split_fingerprint_keys(raw_keys)
        claims.append({
            "claim_id": str(claim_id).strip(),
            "text": str(text or "").strip(),
            "severity": str(severity or "minor").strip().lower(),
            "references_entities": refs,
            "gold_grounding_status": str(gold_grounding_status or "uncertain_needs_drilldown").strip(),
            "gold_linked_fingerprint_keys": keys,
        })

    return {"sample_id": str(sample_id).strip(), "actors": actors, "claims": claims}


def _split_fingerprint_keys(raw: str) -> list[str]:
    """Splits a comma-separated list of fingerprint_key strings, correctly
    handling the fact that each key is itself a JSON array containing
    commas (e.g. '["ARP",["2","aa:bb"]],["UDP",["53"]]' -> two keys)."""
    if not raw:
        return []
    keys, depth, current = [], 0, ""
    for ch in raw:
        if ch == "[":
            depth += 1
        elif ch == "]":
            depth -= 1
        if ch == "," and depth == 0:
            if current.strip():
                keys.append(current.strip())
            current = ""
        else:
            current += ch
    if current.strip():
        keys.append(current.strip())
    return keys


def main():
    parser = argparse.ArgumentParser(description="Convert a filled gold-sample template to gold_sample.json")
    parser.add_argument("xlsx_path")
    parser.add_argument("-o", "--output", default="gold_sample.json")
    args = parser.parse_args()

    result = convert(args.xlsx_path)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f"Wrote {len(result['actors'])} actor(s) and {len(result['claims'])} claim(s) to: {args.output}")


if __name__ == "__main__":
    main()
