"""
Converts a filled-in gold_sample_template.xlsx into the gold_sample.json
shape run_evaluator.py expects. The example row (row 2 on each sheet) is
always skipped by FIXED POSITION, not by matching its content -- an
earlier version matched on the example's id values ("kali_entity", "c1"),
which would silently drop a real user's own data if they happened to
reuse a similarly plausible id (confirmed directly: this happened
immediately when testing with realistic values).

Evidence column syntax: "fingerprint_key" or "fingerprint_key@dest_ip",
comma-separated for multiple entries. The optional "@dest_ip" suffix
disambiguates WHICH conversation a fingerprint_key refers to, needed
because fingerprint_key is deliberately IP-agnostic and can be shared by
entirely different real-world conversations (confirmed on a real
ransomware capture: a lure fetch, a payload download, and a C2 check-in
all produced the identical fingerprint_key). Omit "@dest_ip" only when
genuinely unambiguous for that capture.

Usage:
    python3 xlsx_to_gold_json.py filled_template.xlsx -o gold_sample.json
"""
from __future__ import annotations

import argparse
import json

import openpyxl


def convert(xlsx_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    ws_info = wb["Info"]
    sample_id = ws_info["B1"].value
    if not sample_id:
        raise ValueError("Info sheet: sample_id (cell B1) is empty -- fill it in first.")

    ws_a = wb["Actors"]
    actors = []
    for row in ws_a.iter_rows(min_row=3, values_only=True):  # row 2 is the fixed example row, always skipped
        entity_id, description_ref_gold, ground_truth_ip = (row + (None, None, None))[:3]
        if not entity_id:
            continue
        # Comma-separated for CDN/load-balanced services with multiple
        # legitimate IPs -- see GoldActor.ground_truth_ips.
        ips = [ip.strip() for ip in str(ground_truth_ip or "").split(",") if ip.strip()]
        actors.append({
            "entity_id": str(entity_id).strip(),
            "description_ref_gold": str(description_ref_gold or "").strip(),
            "ground_truth_ips": ips,
        })

    ws_c = wb["Steps"] if "Steps" in wb.sheetnames else wb["Claims"]
    claims = []
    for row in ws_c.iter_rows(min_row=3, values_only=True):  # row 2 is the fixed example row, always skipped
        (claim_id, text, references_entities,
         gold_grounding_status, gold_linked_evidence_raw) = (row + (None,) * 5)[:5]
        if not claim_id:
            continue
        refs = [e.strip() for e in str(references_entities or "").split(",") if e.strip()]
        raw = str(gold_linked_evidence_raw or "").strip()
        evidence = _parse_evidence_refs(raw)
        claims.append({
            "claim_id": str(claim_id).strip(),
            "text": str(text or "").strip(),
            "references_entities": refs,
            "gold_grounding_status": str(gold_grounding_status or "uncertain_needs_drilldown").strip(),
            "gold_linked_evidence": evidence,
        })

    return {"sample_id": str(sample_id).strip(), "actors": actors, "claims": claims}


def _parse_evidence_refs(raw: str) -> list[dict]:
    """
    Parses "key1@ip1,key2,key3@ip3" into
    [{"fingerprint_key": "key1", "flow_dst_ip": "ip1"},
     {"fingerprint_key": "key2", "flow_dst_ip": None},
     {"fingerprint_key": "key3", "flow_dst_ip": "ip3"}]
    Splitting on "," respects JSON-array brackets inside each
    fingerprint_key (which themselves contain commas) -- same
    bracket-depth-aware logic as before, applied before the "@" split.
    """
    segments = _split_bracket_aware(raw)
    refs = []
    for seg in segments:
        seg = seg.strip()
        if not seg:
            continue
        if "@" in seg:
            # Split on the LAST "@" -- a fingerprint_key itself never
            # contains "@", so this is unambiguous even if a MAC address
            # or other field oddly contained one.
            key, ip = seg.rsplit("@", 1)
            refs.append({"fingerprint_key": key.strip(), "flow_dst_ip": ip.strip() or None})
        else:
            refs.append({"fingerprint_key": seg, "flow_dst_ip": None})
    return refs


def _split_bracket_aware(raw: str) -> list[str]:
    if not raw:
        return []
    parts, depth, current = [], 0, ""
    for ch in raw:
        if ch == "[":
            depth += 1
        elif ch == "]":
            depth -= 1
        if ch == "," and depth == 0:
            if current.strip():
                parts.append(current.strip())
            current = ""
        else:
            current += ch
    if current.strip():
        parts.append(current.strip())
    return parts


def main():
    parser = argparse.ArgumentParser(description="Convert a filled gold-sample template to gold_sample.json")
    parser.add_argument("xlsx_path")
    parser.add_argument("-o", "--output", default="gold_sample.json")
    args = parser.parse_args()

    result = convert(args.xlsx_path)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f"Wrote {len(result['actors'])} actor(s) and {len(result['claims'])} claim(s) to: {args.output}")


if __name__ == "__main__":
    main()
